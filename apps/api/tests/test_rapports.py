from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.test_parcelles import ferme  # noqa: F401
from tests.test_plants import client  # noqa: F401


@pytest.fixture()
def avec_arrosage(ferme):  # noqa: F811
    pid = ferme.get("/parcelles").json()[0]["id"]
    assert ferme.post(f"/parcelles/{pid}/vanne", json={"action": "ouvrir", "duree_min": 30}).status_code in (200, 201)
    return ferme


def test_excel_contient_les_feuilles_et_les_chiffres_de_l_ecran(avec_arrosage):
    r = avec_arrosage.get("/rapports/ferme.xlsx?jours=7")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert r.headers["content-disposition"].startswith('attachment; filename="rapport-ferme-') and ".xlsx" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Résumé", "Eau par parcelle", "Eau par jour", "Arrosages", "Alertes", "Parcelles"]
    ecran = avec_arrosage.get("/eau?jours=7").json()
    parcelles = list(wb["Eau par parcelle"].values)
    assert len(parcelles) == 1 + len(ecran["parcelles"]) + 1               # en-tête + parcelles + total
    assert parcelles[-1][0] == "TOTAL FERME" and abs(parcelles[-1][2] - ecran["ferme"]["reels"]) <= 20   # le temps passe entre les deux appels
    assert len(list(wb["Arrosages"].values)) >= 2                            # l'arrosage déclenché y figure
    assert len(list(wb["Parcelles"].values)) == 1 + len(avec_arrosage.get("/parcelles").json())


def test_pdf_valide(avec_arrosage):
    r = avec_arrosage.get("/rapports/ferme.pdf?jours=1")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF") and len(r.content) > 3000 and ".pdf" in r.headers["content-disposition"]


def test_periode_invalide_et_droits(ferme):  # noqa: F811
    assert ferme.get("/rapports/ferme.pdf?jours=0").status_code == 422
    assert ferme.get("/rapports/ferme.xlsx?jours=500").status_code == 422
    ferme.headers.pop("Authorization", None)
    assert ferme.get("/rapports/ferme.pdf").status_code == 401


def test_rapport_sans_aucune_donnee_ne_plante_pas(client):  # noqa: F811
    c, _ = client
    assert c.get("/rapports/ferme.pdf").status_code == 200
    assert c.get("/rapports/ferme.xlsx").status_code == 200


def test_pdf_supporte_les_caracteres_speciaux_dans_les_noms(ferme):  # noqa: F811
    pid = ferme.get("/parcelles").json()[0]["id"]
    assert ferme.put(f"/parcelles/{pid}", json={"nom": "Parcelle R&D <b>test"}).status_code == 200
    r = ferme.get("/rapports/ferme.pdf?jours=1")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert ferme.get("/rapports/ferme.xlsx?jours=1").status_code == 200


def test_module_manquant_donne_un_message_clair(ferme, monkeypatch):  # noqa: F811
    from app import rapports

    def casse(d):
        raise ImportError("pas là", name="reportlab")
    monkeypatch.setattr(rapports, "en_pdf", casse)
    r = ferme.get("/rapports/ferme.pdf")
    assert r.status_code == 503 and "reportlab" in r.json()["detail"] and "pip install" in r.json()["detail"]


def test_les_dates_du_rapport_sont_reelles_meme_si_la_simulation_est_dans_le_futur(ferme, monkeypatch):  # noqa: F811
    from datetime import datetime, timedelta, timezone

    from app import models_parcelles as mp
    from app import rapports
    from app import simulation as sim
    monkeypatch.setattr(sim, "SECONDES", 5.0)
    futur = datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(days=160, hours=3)
    monkeypatch.setitem(mp._horloge, "simulee", futur)
    pid = ferme.get("/parcelles").json()[0]["id"]
    assert ferme.post(f"/parcelles/{pid}/vanne", json={"action": "ouvrir", "duree_min": 30}).status_code in (200, 201)
    from app.database import get_db
    from app.main import app
    db = next(app.dependency_overrides[get_db]())
    d = rapports.donnees(db, 1)
    an = str(datetime.now().year)
    assert d["decalage_jours"] == 160
    assert d["debut"][6:10] == an and d["fin"][6:10] == an and d["genere"][6:10] == an
    assert d["journal"] and all(j["quand"][6:10] == an for j in d["journal"])
    assert all(x["date"][:4] == an for x in d["bilan"]["par_jour"])
    xlsx = ferme.get("/rapports/ferme.xlsx?jours=1")
    assert xlsx.status_code == 200 and ferme.get("/rapports/ferme.pdf?jours=1").status_code == 200
