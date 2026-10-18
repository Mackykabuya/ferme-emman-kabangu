"""Rapports de la ferme en Excel (.xlsx) et en PDF : eau consommée (réelle et « classique »), journal des arrosages,
alertes en cours, état des parcelles. Les chiffres viennent des mêmes calculs que l'écran (aucun recalcul à part)."""
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from sqlalchemy import select
from sqlalchemy.orm import Session

from . import simulation
from .engine import eau as moteur_eau
from .models import Plant
from .models_parcelles import Arrosage, Parcelle, maintenant
from .routers.alertes import toutes_les_alertes
from .routers.eau import consommation

LOGO = Path(__file__).parent / "static" / "logo.png"
GRAVITE = {"critique": "Urgent", "attention": "À vérifier", "info": "Information"}
MESURES = [("humidite_sol_pct", "Humidité sol (%)"), ("temperature_sol_c", "Temp. sol (°C)"), ("ph", "pH"),
           ("azote_N_mgkg", "Azote N"), ("phosphore_P_mgkg", "Phosphore P"), ("potassium_K_mgkg", "Potassium K")]
MAX_LIGNES_JOURNAL = 5000


def _local(iso_ou_date) -> datetime:
    d = datetime.fromisoformat(iso_ou_date) if isinstance(iso_ou_date, str) else iso_ou_date
    return d.replace(tzinfo=timezone.utc).astimezone(moteur_eau.fuseau())


def _fmt(d: datetime) -> str:
    return d.strftime("%d/%m/%Y %H:%M")


def donnees(db: Session, jours: int) -> dict:
    bilan = consommation(jours, db)
    debut, fin = datetime.fromisoformat(bilan["debut"]), datetime.fromisoformat(bilan["fin"])
    parcelles = list(db.scalars(select(Parcelle).order_by(Parcelle.nom)))
    noms = {p.id: p.nom for p in parcelles}
    plantes = {p.id: db.get(Plant, p.plant_id).nom for p in parcelles}
    journal = [{"quand": _fmt(_local(a.at)), "parcelle": noms.get(a.parcelle_id, "?"), "culture": plantes.get(a.parcelle_id, ""),
                "minutes": a.minutes, "litres": round(a.litres), "source": "Automatique" if a.source == "auto" else "Manuel"}
               for a in db.scalars(select(Arrosage).where(Arrosage.action == "ouvrir", Arrosage.at >= debut, Arrosage.at <= fin)
                                   .order_by(Arrosage.at).limit(MAX_LIGNES_JOURNAL))]
    cultures = {p["id"]: plantes[p["id"]] for p in bilan["parcelles"]}
    return {
        "jours": jours, "bilan": bilan, "debut": _fmt(_local(debut)), "fin": _fmt(_local(fin)),
        "genere": _fmt(_local(maintenant())), "simule": simulation.SECONDES > 0, "cultures": cultures, "journal": journal,
        "alertes": [{"gravite": GRAVITE[a["gravite"]], "parcelle": a["parcelle"], "titre": a["titre"], "message": a["message"],
                     "conseil": a["conseil"]} for a in toutes_les_alertes(db)],
        "etat": [{"nom": p.nom, "culture": plantes[p.id], "surface_m2": round(p.surface_m2), "sol": p.type_sol, "stade": p.stade,
                  "mode": "Automatique" if p.mode == "auto" else "Manuel", **{c: p.mesures.get(c) for c, _ in MESURES}}
                 for p in parcelles],
    }


def nom_fichier(d: dict, ext: str) -> str:
    return f"rapport-ferme-{datetime.now().strftime('%Y-%m-%d')}-{d['jours']}j.{ext}"


class Brut(str):
    """Texte déjà mis en forme pour le PDF (balises voulues) : il n'est pas échappé."""


def _pdf(x) -> str:
    """Texte venant des données (noms, alertes…) : « & », « < » ne doivent jamais être pris pour du formatage."""
    return x if isinstance(x, Brut) else ("—" if x is None else escape(str(x)))


def _pct(x):
    return "—" if x is None else f"{x} %"


# ---------------------------------------------------------------- Excel
def en_xlsx(d: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb, b = Workbook(), d["bilan"]
    tete = PatternFill("solid", fgColor="17301F")

    def feuille(titre, colonnes, lignes, largeurs, premiere=False):
        ws = wb.active if premiere else wb.create_sheet()
        ws.title = titre
        ws.append(colonnes)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), tete, Alignment(vertical="center", wrap_text=True)
        for l in lignes:
            ws.append(list(l))
        for i, w in enumerate(largeurs, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    f = b["ferme"]
    r = feuille("Résumé", ["Rapport", "Valeur"], [
        ("Ferme", "Ferme Emman Kabangu"), ("Période", f"du {d['debut']} au {d['fin']} ({str(b['jours']).replace('.', ',')} jour(s))"),
        ("Rapport généré le", d["genere"]),
        ("Eau réellement envoyée (L)", f["reels"]), ("Arrosage classique équivalent (L)", f["classiques"]),
        ("Eau économisée (L)", f["economie"]), ("Économie (%)", f["economie_pct"] if f["economie_pct"] is not None else "—"),
        ("Alertes en cours", len(d["alertes"])), ("Arrosages dans la période", len(d["journal"])),
        ("Méthode de comparaison", b["methode"]),
        ("Remarque", "Comparaison peu fiable : moins d'un jour de données." if not b["comparaison_fiable"] else "")]
        + ([("Données", "SIMULÉES : capteurs et dates simulés, le temps passe plus vite que la réalité.")] if d["simule"] else []),
        [34, 90], premiere=True)
    r["B10"].alignment = Alignment(wrap_text=True, vertical="top")
    feuille("Eau par parcelle", ["Parcelle", "Culture", "Eau réelle (L)", "Classique (L)", "Économie (L)", "Économie (%)"],
            [(p["nom"], d["cultures"].get(p["id"], ""), p["reels"], p["classiques"], p["economie"], p["economie_pct"] if p["economie_pct"] is not None else "—")
             for p in b["parcelles"]] + [("TOTAL FERME", "", f["reels"], f["classiques"], f["economie"], f["economie_pct"] if f["economie_pct"] is not None else "—")],
            [22, 20, 16, 16, 16, 14])
    wb["Eau par parcelle"][f"A{len(b['parcelles']) + 2}"].font = Font(bold=True)
    feuille("Eau par jour", ["Date", "Eau réelle (L)", "Classique (L)", "Économie (L)"],
            [(_jour(x["date"]), x["reels"], x["classiques"], x["classiques"] - x["reels"]) for x in b["par_jour"]], [14, 16, 16, 16])
    feuille("Arrosages", ["Date et heure", "Parcelle", "Culture", "Durée (min)", "Eau prévue (L)", "Déclenché par"],
            [(j["quand"], j["parcelle"], j["culture"], j["minutes"], j["litres"], j["source"]) for j in d["journal"]], [18, 20, 18, 12, 12, 16])
    ws = feuille("Alertes", ["Gravité", "Parcelle", "Problème", "Explication", "À faire"],
                 [(a["gravite"], a["parcelle"], a["titre"], a["message"], a["conseil"]) for a in d["alertes"]] or [("Aucune alerte en cours", "", "", "", "")],
                 [14, 20, 30, 60, 60])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    feuille("Parcelles", ["Parcelle", "Culture", "Surface (m²)", "Sol", "Stade", "Mode"] + [n for _, n in MESURES],
            [(e["nom"], e["culture"], e["surface_m2"], e["sol"], e["stade"], e["mode"], *[e[c] for c, _ in MESURES]) for e in d["etat"]],
            [20, 18, 13, 12, 24, 14, 16, 14, 8, 12, 14, 14])
    sortie = BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


def _jour(iso: str) -> str:
    y, m, j = iso.split("-")
    return f"{j}/{m}/{y}"


# ---------------------------------------------------------------- PDF
def en_pdf(d: dict) -> bytes:
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    VERT, GRIS = colors.HexColor("#17301F"), colors.HexColor("#5D6F62")
    st = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=st["Title"], fontSize=20, textColor=VERT, alignment=0, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=st["Heading2"], fontSize=13, textColor=VERT, spaceBefore=14, spaceAfter=6)
    txt = ParagraphStyle("txt", parent=st["BodyText"], fontSize=9, leading=12)
    petit = ParagraphStyle("petit", parent=txt, fontSize=8, textColor=GRIS)
    cell = ParagraphStyle("cell", parent=txt, fontSize=8, leading=10)

    def tableau(entetes, lignes, largeurs, droite=()):
        t = Table([[Paragraph(f"<b>{h}</b>", ParagraphStyle("th", parent=cell, textColor=colors.white)) for h in entetes]]
                  + [[Paragraph(_pdf(c), cell) for c in l] for l in lignes], colWidths=[w * cm for w in largeurs], repeatRows=1)
        s = [("BACKGROUND", (0, 0), (-1, 0), VERT), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5EF")]),
             ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#D5DDD3")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        s += [("ALIGN", (c, 0), (c, -1), "RIGHT") for c in droite]
        t.setStyle(TableStyle(s))
        return t

    b, f = d["bilan"], d["bilan"]["ferme"]
    n = lambda x: f"{x:,}".replace(",", " ")
    el = []
    entete = [Paragraph("Rapport de la ferme", h1), Paragraph("Ferme Emman Kabangu", txt),
              Paragraph(f"Période : du {d['debut']} au {d['fin']} ({str(b['jours']).replace('.', ',')} jour(s)) · généré le {d['genere']}", petit)]
    if LOGO.exists():
        el.append(Table([[Image(str(LOGO), width=2.2 * cm, height=2.2 * cm, kind="proportional"), entete]], colWidths=[2.8 * cm, 14 * cm],
                        style=[("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    else:
        el += entete
    if d["simule"]:
        el += [Spacer(1, 6), Paragraph("<b>Données simulées :</b> les capteurs et les dates sont simulés (le temps passe plus vite que la réalité).", petit)]
    el.append(Spacer(1, 10))
    tuiles = Table([[Paragraph(f"<font size=15><b>{n(f['reels'])}</b></font><br/>litres envoyés", txt),
                     Paragraph(f"<font size=15><b>{n(f['classiques'])}</b></font><br/>litres en arrosage classique", txt),
                     Paragraph(f"<font size=15><b>{n(f['economie'])}</b></font><br/>litres économisés ({_pct(f['economie_pct'])})", txt),
                     Paragraph(f"<font size=15><b>{len(d['alertes'])}</b></font><br/>alerte(s) en cours", txt)]], colWidths=[4.3 * cm] * 4,
                   style=[("BOX", (0, 0), (-1, -1), .5, colors.HexColor("#D5DDD3")), ("INNERGRID", (0, 0), (-1, -1), .25, colors.HexColor("#D5DDD3")),
                          ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F5EF")), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)])
    el += [tuiles, Spacer(1, 4), Paragraph(b["methode"], petit)]
    if not b["comparaison_fiable"]:
        el.append(Paragraph("<b>Attention :</b> moins d'un jour de données, la comparaison n'est pas encore fiable.", petit))

    el.append(Paragraph("Eau par parcelle", h2))
    lignes = [(p["nom"], d["cultures"].get(p["id"], ""), n(p["reels"]), n(p["classiques"]), n(p["economie"]), _pct(p["economie_pct"])) for p in b["parcelles"]]
    lignes.append(tuple(Brut(x) for x in ("<b>Total ferme</b>", "", f"<b>{n(f['reels'])}</b>", f"<b>{n(f['classiques'])}</b>", f"<b>{n(f['economie'])}</b>", f"<b>{_pct(f['economie_pct'])}</b>")))
    el.append(tableau(["Parcelle", "Culture", "Réelle (L)", "Classique (L)", "Économie (L)", "Économie"], lignes, [3.4, 3.6, 2.6, 2.8, 2.8, 2.2], droite=(2, 3, 4, 5)))

    if b["par_jour"]:
        dess = Drawing(480, 170)
        ch = VerticalBarChart()
        ch.x, ch.y, ch.width, ch.height = 45, 30, 420, 115
        jours = b["par_jour"][-31:]
        ch.data = [[x["reels"] for x in jours], [x["classiques"] for x in jours]]
        ch.categoryAxis.categoryNames = [_jour(x["date"])[:5] for x in jours]
        ch.categoryAxis.labels.fontSize, ch.categoryAxis.labels.angle, ch.categoryAxis.labels.dy = 6, 45, -8
        ch.valueAxis.labels.fontSize, ch.valueAxis.valueMin = 7, 0
        ch.bars[0].fillColor, ch.bars[1].fillColor = colors.HexColor("#1D6FA5"), colors.HexColor("#B7B7B7")
        ch.barSpacing, ch.groupSpacing = 0, 3
        ch.bars.strokeColor = None
        ch.valueAxis.labels.fontName = ch.categoryAxis.labels.fontName = "Helvetica"
        dess.add(ch)
        dess.add(String(45, 155, "Eau par jour (litres)", fontSize=8, fontName="Helvetica", fillColor=GRIS))
        dess.add(String(300, 155, "Bleu : réelle · Gris : arrosage classique", fontSize=7, fontName="Helvetica", fillColor=GRIS))
        el += [Paragraph("Eau par jour", h2), dess]

    el.append(Paragraph(f"Journal des arrosages ({len(d['journal'])})", h2))
    if d["journal"]:
        el.append(tableau(["Date et heure", "Parcelle", "Culture", "Durée (min)", "Eau prévue (L)", "Déclenché par"],
                          [(j["quand"], j["parcelle"], j["culture"], j["minutes"], n(j["litres"]), j["source"]) for j in d["journal"]],
                          [3.3, 3.3, 3.3, 2.2, 2.2, 3.1], droite=(3, 4)))
    else:
        el.append(Paragraph("Aucun arrosage sur la période.", txt))

    el.append(Paragraph("Alertes en cours", h2))
    if d["alertes"]:
        el.append(tableau(["Gravité", "Parcelle", "Problème", "À faire"], [(a["gravite"], a["parcelle"], a["titre"], a["conseil"]) for a in d["alertes"]],
                          [2.3, 3, 4.6, 7.5]))
    else:
        el.append(Paragraph("Aucune alerte en cours.", txt))

    el.append(Paragraph("État des parcelles", h2))
    el.append(KeepTogether(tableau(["Parcelle", "Culture", "Surface (m²)", "Sol", "Humid. (%)", "pH", "N", "P", "K"],
                                   [(e["nom"], e["culture"], n(e["surface_m2"]), e["sol"], e["humidite_sol_pct"], e["ph"], e["azote_N_mgkg"],
                                     e["phosphore_P_mgkg"], e["potassium_K_mgkg"]) for e in d["etat"]],
                                   [2.9, 3.2, 2.1, 2, 1.9, 1.2, 1.2, 1.2, 1.3], droite=(2, 4, 5, 6, 7, 8))))

    def pied(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(GRIS)
        canvas.drawString(1.8 * cm, 1.1 * cm, "Ferme Emman Kabangu · rapport généré automatiquement")
        canvas.drawRightString(A4[0] - 1.8 * cm, 1.1 * cm, f"Page {doc.page}")
        canvas.restoreState()

    sortie = BytesIO()
    SimpleDocTemplate(sortie, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.6 * cm, bottomMargin=1.8 * cm,
                      title="Rapport de la ferme", author="Ferme Emman Kabangu").build(el, onFirstPage=pied, onLaterPages=pied)
    return sortie.getvalue()
